#!/usr/bin/env python3
"""
Excel (.xlsx) → Word (.docx) хөрвүүлэгч.

Хүснэгтийн загварыг алдагдуулахгүйгээр хөрвүүлнэ:
  • нэгтгэсэн нүд (merge), баганын өргөн, мөрийн өндөр
  • хүрээ (border), дэвсгэр өнгө, фонт (нэр, хэмжээ, тод, налуу, доогуур зураас, өнгө)
  • нүд доторх хэсэгчилсэн форматтай текст (rich text)
  • хэвтээ/босоо зэрэгцүүлэлт, догол, босоо (эргүүлсэн) бичвэр
  • тооны формат (0.00, #,##0, %, огноо) — Excel-д харагдах утгаар
  • нуусан мөр/багана алгасна, хэвлэх муж (print area), гарчиг мөрийн давталт
  • хуудасны чиглэл (хэвтээ/босоо), цаасны хэмжээ, захын зай
  • хүснэгт хуудсанд багтахгүй бол Excel-ийн "Fit to page" шиг жигд жижигрүүлнэ

Ашиглах:
  python excel2word.py                      # цонхтой (GUI) горим
  python excel2word.py файл.xlsx            # файл.docx үүснэ
  python excel2word.py файл.xlsx -o гаралт.docx --sheets "Хуудас1" "Хуудас2"
"""

import argparse
import colorsys
import datetime as dt
import os
import re
import sys
from copy import deepcopy
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

try:
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.styles.colors import COLOR_INDEX
    from openpyxl.utils import range_boundaries
    import docx
    from docx.oxml import parse_xml
except ImportError:  # pragma: no cover
    print("Шаардлагатай сан алга. Дараах командыг ажиллуулна уу:\n"
          "    pip install openpyxl python-docx")
    sys.exit(1)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSDECL = f'xmlns:w="{W_NS}"'

# Excel-ийн paperSize код → (өргөн, өндөр) twip-ээр (1 инч = 1440 twip)
PAPER_SIZES = {
    1: (12240, 15840),   # Letter
    5: (12240, 20160),   # Legal
    8: (16838, 23811),   # A3
    9: (11906, 16838),   # A4
    11: (8391, 11906),   # A5
    13: (10318, 14570),  # B5
}

# Excel border style → (Word val, хэмжээ 1/8 pt)
BORDER_MAP = {
    "hair": ("single", 2),
    "thin": ("single", 4),
    "medium": ("single", 12),
    "thick": ("single", 18),
    "double": ("double", 4),
    "dotted": ("dotted", 4),
    "dashed": ("dashed", 4),
    "mediumDashed": ("dashed", 12),
    "dashDot": ("dotDash", 4),
    "mediumDashDot": ("dotDash", 12),
    "dashDotDot": ("dotDotDash", 4),
    "mediumDashDotDot": ("dotDotDash", 12),
    "slantDashDot": ("dotDash", 12),
}
BORDER_WEIGHT = {"hair": 1, "dotted": 2, "dashDotDot": 3, "dashDot": 4,
                 "dashed": 5, "thin": 6, "mediumDashDotDot": 7,
                 "slantDashDot": 8, "mediumDashDot": 9, "mediumDashed": 10,
                 "medium": 11, "thick": 12, "double": 13}

H_ALIGN = {"left": "left", "center": "center", "centerContinuous": "center",
           "right": "right", "justify": "both", "distributed": "distribute",
           "fill": "left"}
V_ALIGN = {"top": "top", "center": "center", "bottom": "bottom",
           "justify": "center", "distributed": "center"}

_INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ----------------------------------------------------------------- өнгө ---
class ColorResolver:
    """openpyxl-ийн Color объектыг 'RRGGBB' болгоно (theme, indexed, tint)."""

    THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                   "accent4", "accent5", "accent6", "hlink", "folHlink"]

    def __init__(self, wb):
        self.theme = self._parse_theme(getattr(wb, "loaded_theme", None))

    def _parse_theme(self, data):
        defaults = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31",
                    "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]
        if not data:
            return defaults
        try:
            root = ET.fromstring(data)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            scheme = root.find(".//a:clrScheme", ns)
            colors = {}
            for el in scheme:
                tag = el.tag.split("}")[1]
                child = el[0]
                val = child.get("val") if child.tag.endswith("srgbClr") \
                    else child.get("lastClr")
                colors[tag] = val
            return [colors.get(k, d) for k, d in zip(self.THEME_ORDER, defaults)]
        except Exception:
            return defaults

    @staticmethod
    def _apply_tint(hex6, tint):
        if not tint:
            return hex6
        r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
        r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
        return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))

    def resolve(self, color):
        if color is None:
            return None
        try:
            ctype = color.type
            if ctype == "rgb":
                val = color.rgb
                if not isinstance(val, str):
                    return None
                hex6 = val[-6:].upper()
            elif ctype == "theme":
                idx = color.theme
                if idx is None or idx >= len(self.theme):
                    return None
                hex6 = self.theme[idx].upper()
            elif ctype == "indexed":
                idx = color.indexed
                if idx is None or idx >= 64:   # 64 = system foreground/auto
                    return None
                hex6 = COLOR_INDEX[idx][-6:].upper()
            else:
                return None
            return self._apply_tint(hex6, color.tint)
        except Exception:
            return None


# --------------------------------------------------------- тооны формат ---
def _fmt_general(v):
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        s = "%.10g" % v
        return s
    return str(v)


def _fmt_date(v, fmt):
    f = re.sub(r"\[[^\]]*\]", "", fmt)
    f = re.sub(r'"([^"]*)"', r"\1", f).replace("\\", "")
    if not re.search(r"[yYmMdDhHsS]", f):
        return v.strftime("%Y-%m-%d")
    out = []
    tokens = re.findall(r"yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|am/pm|.",
                        f, flags=re.I)
    prev_h = False
    for idx, t in enumerate(tokens):
        tl = t.lower()
        nxt = "".join(tokens[idx + 1:idx + 3]).lower()
        if tl == "yyyy":
            out.append("%04d" % v.year)
        elif tl == "yy":
            out.append("%02d" % (v.year % 100))
        elif tl in ("mm", "m"):
            is_min = prev_h or nxt.startswith(":s") or nxt.startswith("s")
            if is_min and isinstance(v, dt.datetime):
                out.append("%02d" % v.minute if tl == "mm" else str(v.minute))
            else:
                out.append("%02d" % v.month if tl == "mm" else str(v.month))
        elif tl == "mmm":
            out.append(v.strftime("%b"))
        elif tl == "mmmm":
            out.append(v.strftime("%B"))
        elif tl == "dd":
            out.append("%02d" % v.day)
        elif tl == "d":
            out.append(str(v.day))
        elif tl in ("ddd", "dddd"):
            out.append(v.strftime("%a" if tl == "ddd" else "%A"))
        elif tl in ("hh", "h"):
            hour = getattr(v, "hour", 0)
            if "am/pm" in f.lower():
                hour = hour % 12 or 12
            out.append("%02d" % hour if tl == "hh" else str(hour))
        elif tl in ("ss", "s"):
            sec = getattr(v, "second", 0)
            out.append("%02d" % sec if tl == "ss" else str(sec))
        elif tl == "am/pm":
            out.append("AM" if getattr(v, "hour", 0) < 12 else "PM")
        else:
            out.append(t)
        if tl in ("hh", "h"):
            prev_h = True
        elif tl != ":":
            prev_h = False
    return "".join(out)


def format_value(v, fmt):
    """Excel-д харагдах байдлаар утгыг текст болгоно."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        if isinstance(v, dt.time):
            return v.strftime("%H:%M:%S" if v.second else "%H:%M")
        return _fmt_date(v, fmt or "yyyy-mm-dd")
    if isinstance(v, dt.timedelta):
        return str(v)
    if not isinstance(v, (int, float)):
        return str(v)
    if not fmt or fmt == "General" or fmt == "@":
        return _fmt_general(v)

    sections = fmt.split(";")
    sec = sections[0]
    neg_section = False
    if v < 0 and len(sections) > 1:
        sec = sections[1]
        neg_section = True
    elif v == 0 and len(sections) > 2:
        sec = sections[2]
    if not sec.strip():          # жишээ нь "0;-0;;@" → тэгийг харуулахгүй
        return ""
    sec = re.sub(r"\[[^\]]*\]", "", sec)          # [Red], [$₮-450] гэх мэт
    sec = re.sub(r"_.", " ", sec).replace("*", "")

    m = re.search(r"[#0?][#0?,]*(\.[#0?]+)?(E[+-]0+)?", sec, flags=re.I)
    if not m:
        return re.sub(r'"([^"]*)"', r"\1", sec).replace("\\", "") or _fmt_general(v)
    num_fmt = m.group(0)
    decimals = len(m.group(1)) - 1 if m.group(1) else 0
    min_dec = m.group(1).count("0") if m.group(1) else 0
    val = abs(v) if neg_section else v
    if "%" in sec:
        val = val * 100
    if m.group(2):
        body = f"{val:.{decimals}E}"
    else:
        thousands = "," in num_fmt
        body = f"{val:,.{decimals}f}" if thousands else f"{val:.{decimals}f}"
        if decimals > min_dec and "." in body:      # '#' оронгууд
            intp, frac = body.split(".")
            frac = frac.rstrip("0")
            frac = frac + "0" * max(0, min_dec - len(frac))
            body = intp + ("." + frac if frac else "")
        int_zeros = num_fmt.split(".")[0].count("0")
        if int_zeros == 0 and body.lstrip("-").startswith("0") and "." in body:
            body = body.replace("0.", ".", 1)
    prefix, suffix = sec[:m.start()], sec[m.end():]
    clean = lambda s: re.sub(r'"([^"]*)"', r"\1", s).replace("\\", "")
    return clean(prefix) + body + clean(suffix)


# --------------------------------------------------------------- хөрвүүлэх ---
def col_width_to_twips(width):
    px = int(width * 7 + 5) if width > 0 else 0   # Calibri 11-д ойролцоо
    return int(px * 0.75 * 20)


def xml_text(s):
    return escape(_INVALID_XML.sub("", s))


class SheetConverter:
    def __init__(self, ws, ws_formulas, colors, fit=True, scale_fonts=True):
        self.ws = ws
        self.wsf = ws_formulas
        self.colors = colors
        self.fit = fit
        self.scale_fonts = scale_fonts
        self.missing_formula_values = 0

    # ---- муж, хэмжээ -------------------------------------------------------
    def _used_range(self):
        ws = self.ws
        if ws.print_area:
            area = ws.print_area
            area = area[0] if isinstance(area, (list, tuple)) else area.split(",")[0]
            ref = area.split("!")[-1].replace("$", "")
            min_c, min_r, max_c, max_r = range_boundaries(ref)
            return min_r, min_c, max_r, max_c

        max_r = max_c = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, "") or self._has_visual(c):
                    max_r = max(max_r, c.row)
                    max_c = max(max_c, c.column)
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= max_r and mr.min_col <= max_c:
                max_r = max(max_r, mr.max_row)
                max_c = max(max_c, mr.max_col)
        return 1, 1, max(max_r, 1), max(max_c, 1)

    @staticmethod
    def _has_visual(c):
        if not c.has_style:
            return False
        b = c.border
        if any(getattr(b, s) is not None and getattr(b, s).style
               for s in ("left", "right", "top", "bottom")):
            return True
        return c.fill is not None and c.fill.fill_type not in (None, "none")

    def _column_info(self, min_c, max_c):
        ws = self.ws
        default = ws.sheet_format.defaultColWidth
        if not default:
            base = ws.sheet_format.baseColWidth or 8
            default = base + 0.71
        widths = {c: default for c in range(min_c, max_c + 1)}
        hidden = set()
        for dim in ws.column_dimensions.values():
            lo, hi = dim.min or 0, dim.max or 0
            for c in range(max(lo, min_c), min(hi, max_c) + 1):
                if dim.hidden:
                    hidden.add(c)
                if dim.width and (dim.customWidth or dim.width != 13):
                    widths[c] = dim.width
        for c in range(min_c, max_c + 1):
            if widths[c] <= 0:
                hidden.add(c)
        return widths, hidden

    # ---- хүрээ ---------------------------------------------------------------
    def _side(self, r, c, side):
        if r < 1 or c < 1:
            return None
        cell = self.ws._cells.get((r, c))
        if cell is None or not cell.has_style:
            return None
        s = getattr(cell.border, side)
        if s is None or not s.style:
            return None
        return s

    def _edge(self, r, c, side):
        """Хөрш нүдний эсрэг талын хүрээг ч тооцож, хүчтэйг нь сонгоно."""
        opposite = {"left": ("right", 0, -1), "right": ("left", 0, 1),
                    "top": ("bottom", -1, 0), "bottom": ("top", 1, 0)}
        own = self._side(r, c, side)
        opp_side, dr, dc = opposite[side]
        nb = self._side(r + dr, c + dc, opp_side)
        cands = [s for s in (own, nb) if s is not None]
        if not cands:
            return None
        return max(cands, key=lambda s: BORDER_WEIGHT.get(s.style, 0))

    def _border_xml(self, tag, side_obj):
        if side_obj is None:
            return f'<w:{tag} w:val="nil"/>'
        val, sz = BORDER_MAP.get(side_obj.style, ("single", 4))
        color = self.colors.resolve(side_obj.color) or "000000"
        return f'<w:{tag} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'

    # ---- текст ---------------------------------------------------------------
    def _rpr(self, font, scale, override=None):
        """font нь openpyxl Font; override нь InlineFont (rich text)."""
        def pick(attr, inline_attr=None):
            if override is not None:
                v = getattr(override, inline_attr or attr, None)
                if v is not None:
                    return v
            return getattr(font, attr, None) if font is not None else None

        parts = []
        name = pick("name", "rFont") or "Calibri"
        name = xml_text(str(name)).replace('"', "&quot;")
        parts.append(f'<w:rFonts w:ascii="{name}" w:hAnsi="{name}" '
                     f'w:cs="{name}" w:eastAsia="{name}"/>')
        if pick("b"):
            parts.append("<w:b/><w:bCs/>")
        if pick("i"):
            parts.append("<w:i/><w:iCs/>")
        if pick("strike"):
            parts.append("<w:strike/>")
        color = None
        if override is not None and getattr(override, "color", None) is not None:
            color = self.colors.resolve(override.color)
        elif font is not None and font.color is not None:
            color = self.colors.resolve(font.color)
        if color:
            parts.append(f'<w:color w:val="{color}"/>')
        size = pick("sz") or pick("size") or 11
        try:
            size = float(size)
        except (TypeError, ValueError):
            size = 11
        if self.scale_fonts:
            size *= scale
        half = max(2, int(round(size * 2)))
        parts.append(f'<w:sz w:val="{half}"/><w:szCs w:val="{half}"/>')
        u = pick("u")
        if u and u != "none":
            parts.append(f'<w:u w:val="{"double" if "double" in str(u) else "single"}"/>')
        va = pick("vertAlign")
        if va in ("superscript", "subscript"):
            parts.append(f'<w:vertAlign w:val="{va}"/>')
        return "<w:rPr>" + "".join(parts) + "</w:rPr>"

    @staticmethod
    def _runs_from_text(text, rpr):
        out = []
        lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        for i, line in enumerate(lines):
            if i:
                out.append(f"<w:r>{rpr}<w:br/></w:r>")
            if line:
                out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{xml_text(line)}</w:t></w:r>')
        return "".join(out)

    def _cell_content(self, cell, scale):
        font = cell.font if cell is not None else None
        base_rpr = self._rpr(font, scale)
        al = cell.alignment if cell is not None else None
        value = cell.value if cell is not None else None

        if value is None and cell is not None and self.wsf is not None:
            fcell = self.wsf._cells.get((cell.row, cell.column))
            if fcell is not None and isinstance(fcell.value, str) \
                    and fcell.value.startswith("="):
                self.missing_formula_values += 1

        runs = ""
        is_number = isinstance(value, (int, float)) and not isinstance(value, bool)
        if isinstance(value, CellRichText):
            for part in value:
                if isinstance(part, TextBlock):
                    runs += self._runs_from_text(part.text, self._rpr(font, scale, part.font))
                else:
                    runs += self._runs_from_text(str(part), base_rpr)
        elif value is not None:
            text = format_value(value, cell.number_format)
            if is_number and value == 0 and self.ws.sheet_view.showZeros is False:
                text = ""
            runs = self._runs_from_text(text, base_rpr)

        # Paragraph шинж
        h = al.horizontal if al is not None else None
        if h in (None, "general"):
            jc = "right" if is_number or isinstance(value, (dt.date, dt.datetime)) \
                else "center" if isinstance(value, bool) else "left"
        else:
            jc = H_ALIGN.get(h, "left")
        ind = ""
        if al is not None and al.indent:
            tw = int(al.indent * 9 * 20 * (scale if self.scale_fonts else 1))
            ind = f'<w:ind w:{"right" if jc == "right" else "left"}="{tw}"/>'
        ppr = (f'<w:pPr><w:spacing w:before="0" w:after="0" w:line="240" '
               f'w:lineRule="auto"/>{ind}<w:jc w:val="{jc}"/>{base_rpr}</w:pPr>')
        return f"<w:p>{ppr}{runs}</w:p>"

    # ---- хүснэгт -------------------------------------------------------------
    def build(self, usable_width, usable_height):
        ws = self.ws
        min_r, min_c, max_r, max_c = self._used_range()
        widths, hidden_cols = self._column_info(min_c, max_c)
        cols = [c for c in range(min_c, max_c + 1) if c not in hidden_cols]
        rows = [r for r in range(min_r, max_r + 1)
                if not (ws.row_dimensions[r].hidden if r in ws.row_dimensions else False)]
        if not cols or not rows:
            return [], 1.0

        col_tw = {c: col_width_to_twips(widths[c]) for c in cols}
        total = sum(col_tw.values())
        default_h = ws.sheet_format.defaultRowHeight or 15
        row_h = {}
        for r in rows:
            rd = ws.row_dimensions[r] if r in ws.row_dimensions else None
            row_h[r] = rd.height if rd is not None and rd.height else default_h

        scale = 1.0
        ps = ws.page_setup
        fit_to_page = bool(ws.sheet_properties.pageSetUpPr
                           and ws.sheet_properties.pageSetUpPr.fitToPage)
        if not fit_to_page and ps.scale and 10 <= int(ps.scale) < 100:
            scale = int(ps.scale) / 100           # Excel-ийн "Adjust to NN%"
        if self.fit and total * scale > usable_width:
            scale = usable_width / total
        if self.fit and fit_to_page:
            pages_h = ps.fitToHeight if ps.fitToHeight is not None else 1
            if pages_h:
                total_h = sum(row_h.values()) * 20 * 1.03   # бага зэрэг нөөц
                scale = min(scale, usable_height * int(pages_h) / total_h)
        col_tw = {c: max(20, int(w * scale)) for c, w in col_tw.items()}
        total = sum(col_tw.values())
        col_index = {c: i for i, c in enumerate(cols)}
        row_index = {r: i for i, r in enumerate(rows)}

        # Нэгтгэсэн нүднүүд (харагдах мөр/баганын индексээр)
        merge_at = {}      # (row, col) -> ("start", colspan, rowspan_flag, rng) | ("cont", ...) | skip
        for mr in ws.merged_cells.ranges:
            vcols = [c for c in range(mr.min_col, mr.max_col + 1) if c in col_index]
            vrows = [r for r in range(mr.min_row, mr.max_row + 1) if r in row_index]
            if not vcols or not vrows:
                continue
            for r in vrows:
                for c in vcols:
                    merge_at[(r, c)] = ("skip", mr)
                vmerge = None
                if len(vrows) > 1:
                    vmerge = "restart" if r == vrows[0] else "continue"
                merge_at[(r, vcols[0])] = ("start", vcols, vmerge, mr, r == vrows[0])

        title_rows = set()
        if ws.print_title_rows:
            m = re.match(r"\$?(\d+):\$?(\d+)", ws.print_title_rows)
            if m:
                title_rows = set(range(int(m.group(1)), int(m.group(2)) + 1))

        def emit(seg_rows):
            header_ok = True
            xml = [f"<w:tbl {NSDECL}>"]
            jc = "center" if ws.print_options.horizontalCentered else "left"
            xml.append(
                f'<w:tblPr><w:tblW w:w="{total}" w:type="dxa"/><w:jc w:val="{jc}"/>'
                f'<w:tblInd w:w="0" w:type="dxa"/>'
                f'<w:tblLayout w:type="fixed"/>'
                f'<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="{int(30 * max(scale, .5))}" w:type="dxa"/>'
                f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="{int(30 * max(scale, .5))}" w:type="dxa"/></w:tblCellMar>'
                f'<w:tblLook w:val="0000"/></w:tblPr>')
            xml.append("<w:tblGrid>" + "".join(f'<w:gridCol w:w="{col_tw[c]}"/>' for c in cols)
                       + "</w:tblGrid>")

            for r in seg_rows:
                h_tw = max(1, int(row_h[r] * 20 * scale))
                trpr = f'<w:trHeight w:val="{h_tw}" w:hRule="atLeast"/>'
                if r in title_rows and header_ok:
                    trpr += "<w:tblHeader/>"
                else:
                    header_ok = False
                xml.append(f"<w:tr><w:trPr><w:cantSplit/>{trpr}</w:trPr>")

                for c in cols:
                    info = merge_at.get((r, c))
                    if info and info[0] == "skip":
                        continue
                    span_cols, vmerge, top_row_of_merge = [c], None, True
                    if info and info[0] == "start":
                        _, span_cols, vmerge, mr, top_row_of_merge = info
                        anchor = ws._cells.get((mr.min_row, mr.min_col))
                        last_c = span_cols[-1]
                        b_top = self._edge(r, c, "top")
                        b_bottom = self._edge(r, c, "bottom")
                        for cc in span_cols:
                            b_top = b_top or self._edge(r, cc, "top")
                            b_bottom = b_bottom or self._edge(r, cc, "bottom")
                        b_left = self._edge(r, c, "left")
                        b_right = self._edge(r, last_c, "right")
                        # Excel-д хүрээ зөвхөн зарим нүдэнд байж болно → бүх мөрөөс хайна
                        for rr in range(mr.min_row, mr.max_row + 1):
                            b_left = b_left or self._edge(rr, mr.min_col, "left")
                            b_right = b_right or self._edge(rr, mr.max_col, "right")
                        if vmerge == "restart":
                            b_bottom = None
                            for rr in [mr.max_row]:
                                for cc in span_cols:
                                    b_bottom = b_bottom or self._edge(rr, cc, "bottom")
                        style_cell = anchor
                    else:
                        style_cell = ws._cells.get((r, c))
                        b_top = self._edge(r, c, "top")
                        b_bottom = self._edge(r, c, "bottom")
                        b_left = self._edge(r, c, "left")
                        b_right = self._edge(r, c, "right")

                    width = sum(col_tw[cc] for cc in span_cols)
                    tcpr = [f'<w:tcW w:w="{width}" w:type="dxa"/>']
                    if len(span_cols) > 1:
                        tcpr.append(f'<w:gridSpan w:val="{len(span_cols)}"/>')
                    if vmerge:
                        tcpr.append(f'<w:vMerge w:val="{vmerge}"/>' if vmerge == "restart"
                                    else "<w:vMerge/>")
                    tcpr.append("<w:tcBorders>"
                                + self._border_xml("top", b_top)
                                + self._border_xml("left", b_left)
                                + self._border_xml("bottom", b_bottom)
                                + self._border_xml("right", b_right)
                                + "</w:tcBorders>")
                    if style_cell is not None and style_cell.has_style:
                        fill = style_cell.fill
                        if fill is not None and fill.fill_type not in (None, "none"):
                            col = None
                            if fill.fill_type == "solid":
                                col = self.colors.resolve(fill.fgColor) or \
                                    self.colors.resolve(fill.bgColor)
                            elif hasattr(fill, "fgColor"):
                                col = self.colors.resolve(fill.fgColor)
                            elif hasattr(fill, "stop") and fill.stop:
                                col = self.colors.resolve(fill.stop[0].color)
                            if col:
                                tcpr.append(f'<w:shd w:val="clear" w:color="auto" w:fill="{col}"/>')
                        al = style_cell.alignment
                        rot = al.textRotation or 0
                        if rot:
                            if rot == 255 or 91 <= rot <= 180:
                                tcpr.append('<w:textDirection w:val="tbRl"/>')
                            elif rot >= 45:
                                tcpr.append('<w:textDirection w:val="btLr"/>')
                        if al.shrink_to_fit:
                            tcpr.append("<w:tcFitText/>")
                        va = V_ALIGN.get(al.vertical or "bottom", "bottom")
                        tcpr.append(f'<w:vAlign w:val="{va}"/>')
                    else:
                        tcpr.append('<w:vAlign w:val="bottom"/>')

                    content_cell = style_cell if top_row_of_merge else None
                    if content_cell is None and info and info[0] == "start":
                        # vMerge continue нүдэнд текст бичихгүй, гэхдээ фонтын хэмжээ хадгална
                        para = self._cell_content(None, scale)
                    elif content_cell is None and not info:
                        para = self._cell_content(None, scale)
                    else:
                        para = self._cell_content(content_cell, scale)
                    xml.append(f"<w:tc><w:tcPr>{''.join(tcpr)}</w:tcPr>{para}</w:tc>")
                xml.append("</w:tr>")
            xml.append("</w:tbl>")
            return parse_xml("".join(xml))

        # Word-д давтагдах гарчиг мөр хүснэгтийн эхэнд байх ёстой тул
        # гарчгаас өмнөх мөрүүдийг тусдаа хүснэгт болгоно.
        segments = [rows]
        if title_rows:
            first_t = min(title_rows)
            pre = [r for r in rows if r < first_t]
            crosses = any(mr.min_row < first_t <= mr.max_row for mr in ws.merged_cells.ranges)
            if pre and len(pre) < len(rows) and not crosses:
                segments = [pre, [r for r in rows if r >= first_t]]
        return [emit(s) for s in segments], scale


def _page_settings(ws):
    ps = ws.page_setup
    w, h = PAPER_SIZES.get(int(ps.paperSize) if ps.paperSize else 9, PAPER_SIZES[9])
    landscape = ps.orientation == "landscape"
    if landscape:
        w, h = h, w
    pm = ws.page_margins
    to_tw = lambda inch, d: int((inch if inch is not None else d) * 1440)
    m = dict(left=to_tw(pm.left, 0.7), right=to_tw(pm.right, 0.7),
             top=to_tw(pm.top, 0.75), bottom=to_tw(pm.bottom, 0.75),
             header=to_tw(pm.header, 0.3), footer=to_tw(pm.footer, 0.3))
    return w, h, landscape, m


def _sectpr_xml(w, h, landscape, m):
    orient = ' w:orient="landscape"' if landscape else ""
    return (f'<w:sectPr {NSDECL}><w:pgSz w:w="{w}" w:h="{h}"{orient}/>'
            f'<w:pgMar w:top="{m["top"]}" w:right="{m["right"]}" w:bottom="{m["bottom"]}" '
            f'w:left="{m["left"]}" w:header="{m["header"]}" w:footer="{m["footer"]}" w:gutter="0"/>'
            f'<w:cols w:space="720"/></w:sectPr>')


def _has_uncached_formulas(wb, wbf):
    for ws in wbf.worksheets:
        vals = wb[ws.title]._cells
        for key, c in ws._cells.items():
            if isinstance(c.value, str) and c.value.startswith("=") \
                    and key in vals and vals[key].value is None:
                return True
    return False


def _find_soffice():
    import shutil
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(p):
            return p
    return None


def _recalc_with_libreoffice(xlsx_path, log):
    """Томьёоны утга хадгалагдаагүй файлыг LibreOffice байвал дахин бодуулна."""
    import subprocess
    import tempfile
    soffice = _find_soffice()
    if not soffice:
        return None
    outdir = tempfile.mkdtemp(prefix="excel2word_")
    try:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx",
                        "--outdir", outdir, os.path.abspath(xlsx_path)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=180, check=False)
    except Exception:
        return None
    out = os.path.join(outdir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".xlsx")
    if os.path.exists(out):
        log("  i Томьёоны утгыг LibreOffice-оор бодуулав")
        return out
    return None


def convert(xlsx_path, docx_path=None, sheets=None, fit=True, scale_fonts=True,
            log=print):
    if not xlsx_path.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Зөвхөн .xlsx / .xlsm файл дэмжинэ. "
                         "Хуучин .xls файлыг Excel-ээр нээгээд .xlsx болгож хадгална уу.")
    if docx_path is None:
        docx_path = os.path.splitext(xlsx_path)[0] + ".docx"

    wbf = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, rich_text=True)
    if _has_uncached_formulas(wb, wbf):
        recalced = _recalc_with_libreoffice(xlsx_path, log)
        if recalced:
            wb = openpyxl.load_workbook(recalced, data_only=True, rich_text=True)
            wbf = None          # утга бодогдсон тул анхааруулга хэрэггүй
    colors = ColorResolver(wb)

    names = sheets or [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    missing = [n for n in names if n not in wb.sheetnames]
    if missing:
        raise ValueError("Ийм нэртэй хуудас алга: " + ", ".join(missing))

    doc = docx.Document()
    body = doc.element.body
    for p in list(body):
        if p.tag != f"{{{W_NS}}}sectPr":
            body.remove(p)
    final_sect = body.find(f"{{{W_NS}}}sectPr")

    total_missing = 0
    done = 0
    for i, name in enumerate(names):
        ws = wb[name]
        wsf = wbf[name] if wbf is not None else None
        w, h, landscape, m = _page_settings(ws)
        usable = w - m["left"] - m["right"]
        usable_h = h - m["top"] - m["bottom"]
        conv = SheetConverter(ws, wsf, colors, fit=fit, scale_fonts=scale_fonts)
        tables, scale = conv.build(usable, usable_h)
        total_missing += conv.missing_formula_values
        if not tables:
            log(f"  – '{name}' хуудас хоосон тул алгаслаа")
            continue
        for tbl in tables:
            final_sect.addprevious(tbl)
            # Хүснэгтийн араас заавал нэг paragraph хэрэгтэй (жижиг, зайгүй)
            tail = parse_xml(
                f'<w:p {NSDECL}><w:pPr><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'
                f'<w:rPr><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr></w:pPr></w:p>')
            final_sect.addprevious(tail)
        sect = parse_xml(_sectpr_xml(w, h, landscape, m))
        is_last = i == len(names) - 1
        if is_last:
            for child in list(final_sect):
                final_sect.remove(child)
            for child in list(sect):
                final_sect.append(deepcopy(child))
        else:
            tail.find(f"{{{W_NS}}}pPr").append(sect)
        done += 1
        info = f"  ✓ '{name}'"
        if scale < 0.999:
            info += f" (хуудсанд багтаахын тулд {scale * 100:.0f}% болгож жижигрүүлэв)"
        log(info)

    if done == 0:
        raise ValueError("Хөрвүүлэх өгөгдөлтэй хуудас олдсонгүй.")
    # Сүүлийн хуудас хоосон байсан бол өмнөх sectPr-ийг эцсийнх болгоно
    last_tail = final_sect.getprevious()
    inner = last_tail.find(f"{{{W_NS}}}pPr/{{{W_NS}}}sectPr") if last_tail is not None else None
    if inner is not None:
        for child in list(final_sect):
            final_sect.remove(child)
        for child in list(inner):
            final_sect.append(deepcopy(child))
        inner.getparent().remove(inner)

    doc.save(docx_path)
    if total_missing:
        log(f"  ! {total_missing} томьёоны үр дүн файлд хадгалагдаагүй байсан тул хоосон гарлаа. "
            "Excel-ээр нээгээд дахин хадгалсны дараа хөрвүүлнэ үү.")
    return docx_path


# ------------------------------------------------------------------- GUI ---
def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("Excel → Word хөрвүүлэгч")
    root.geometry("620x420")

    fit_var = tk.BooleanVar(value=True)

    frm = tk.Frame(root, padx=12, pady=12)
    frm.pack(fill="both", expand=True)
    tk.Label(frm, text="Excel файлаа сонгоод Word болгон хөрвүүлнэ.\n"
                       "Word файл Excel файлтай ижил хавтсанд үүснэ.",
             justify="left").pack(anchor="w")
    tk.Checkbutton(frm, text="Хүснэгт хуудсанд багтахгүй бол жижигрүүлэх",
                   variable=fit_var).pack(anchor="w", pady=(8, 0))
    out = scrolledtext.ScrolledText(frm, height=14)
    out.pack(fill="both", expand=True, pady=8)

    def log(msg):
        out.insert("end", msg + "\n")
        out.see("end")
        root.update_idletasks()

    def choose():
        files = filedialog.askopenfilenames(
            title="Excel файл сонгох",
            filetypes=[("Excel файл", "*.xlsx *.xlsm"), ("Бүх файл", "*.*")])
        ok = 0
        for f in files:
            log(f"▶ {os.path.basename(f)}")
            try:
                res = convert(f, fit=fit_var.get(), log=log)
                log(f"  → {res}\n")
                ok += 1
            except Exception as e:
                log(f"  ✗ Алдаа: {e}\n")
        if files:
            messagebox.showinfo("Дууслаа", f"{ok}/{len(files)} файл амжилттай хөрвүүлэгдлээ.")

    tk.Button(frm, text="Excel файл сонгох…", command=choose,
              height=2).pack(fill="x")
    root.mainloop()


def main(argv=None):
    ap = argparse.ArgumentParser(description="Excel (.xlsx) файлыг загварыг нь "
                                             "алдагдуулахгүйгээр Word (.docx) болгоно.")
    ap.add_argument("files", nargs="*", help="Excel файл(ууд)")
    ap.add_argument("-o", "--output", help="Гаралтын .docx (ганц файл өгсөн үед)")
    ap.add_argument("--sheets", nargs="+", help="Зөвхөн эдгээр хуудсыг хөрвүүлэх")
    ap.add_argument("--no-fit", action="store_true",
                    help="Хуудсанд багтаахаар жижигрүүлэхгүй")
    ap.add_argument("--keep-font-size", action="store_true",
                    help="Жижигрүүлэхдээ фонтын хэмжээг өөрчлөхгүй (зөвхөн баганын өргөн)")
    args = ap.parse_args(argv)

    if not args.files:
        run_gui()
        return 0
    if args.output and len(args.files) > 1:
        ap.error("-o-г зөвхөн ганц файлтай ашиглана")
    rc = 0
    for f in args.files:
        print(f"▶ {f}")
        try:
            res = convert(f, args.output, sheets=args.sheets, fit=not args.no_fit,
                          scale_fonts=not args.keep_font_size)
            print(f"  → {res}")
        except Exception as e:
            print(f"  ✗ Алдаа: {e}")
            rc = 1
    return rc


if __name__ == "__main__":
    sys.exit(main())
