"""7 хоногийн төлөвлөлт (.xlsx) + чиглүүлэгтэй үйл ажиллагааны өдөр бүрийн төлөвлөлт (.docx)-ийг
PROMPT.md-ийн загвараар нэг Word файл болгоно:
  1) 7 хоногийн төлөвлөлт — хөндлөн хуудас, Excel-ийн хүснэгтээр
  2) 5 өдрийн төлөвлөлт — босоо хуудас, өдөр бүр шинэ хуудаснаас (docx-ийн бичвэр, зураг хэвээр)

Хэрэглээ: python3 week_xlsx_with_days.py 7ХОНОГ.xlsx ӨДРҮҮД.docx ГАРАХ.docx
"""
import copy, os, re, sys, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from lxml import etree
from split_weeks import W, CM, TOP, LEFT, BOTTOM, RIGHT, SIZE, FONT, apply_font, list_sud, set_flag, txt

TABLE_SIZE = 18   # 7 хоногийн хүснэгт 9pt — өмнөх Word загваруудын 7 хоногийн хүснэгттэй адил, хөндлөн хуудсанд багтаана
DAY_START = re.compile(r'^\d+ (дэх|дахь) өдөр\s*:')


def E(tag, **attrs):
    el = etree.Element(W + tag)
    for k, v in attrs.items():
        el.set(W + k, str(v))
    return el


def run(text, bold=False, size=SIZE):
    r = E('r')
    rpr = etree.SubElement(r, W + 'rPr')
    rpr.append(E('rFonts', ascii=FONT, hAnsi=FONT, cs=FONT, eastAsia=FONT))
    if bold:
        rpr.append(E('b'))
    rpr.append(E('sz', val=size)); rpr.append(E('szCs', val=size))
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        if i:
            r.append(E('br'))
        t = etree.SubElement(r, W + 't'); t.text = ln
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return r


def para(parts, jc=None, after=40, keep=False, size=SIZE):
    p = E('p')
    ppr = etree.SubElement(p, W + 'pPr')
    if keep:
        ppr.append(E('keepNext'))
    ppr.append(E('spacing', before=0, after=after, line=252, lineRule='auto'))
    if jc:
        ppr.append(E('jc', val=jc))
    for t, b in ([(parts, False)] if isinstance(parts, str) else parts):
        p.append(run(t, b, size))
    return p


def cell(width, paras, fill=None, span=1):
    tc = E('tc')
    pr = etree.SubElement(tc, W + 'tcPr')
    pr.append(E('tcW', w=width, type='dxa'))
    if span > 1:
        pr.append(E('gridSpan', val=span))
    if fill:
        pr.append(E('shd', val='clear', color='auto', fill=fill))
    for p in paras:
        tc.append(p)
    return tc


def bold_title_paras(title, desc):
    """«Нэр. КОД.» мөрийг тод, тайлбарыг энгийнээр."""
    out = []
    if title:
        out.append(para([(title, True)], after=20, size=TABLE_SIZE))
    for ln in (desc or '').split('\n'):
        if ln.strip():
            m = re.match(r'^([^:]{1,60}):\s*(.*)$', ln)
            out.append(para([(m.group(1) + ':', True), (' ' + m.group(2), False)] if m and m.group(2)
                            else [(ln, ln.rstrip().endswith(':'))], after=20, size=TABLE_SIZE))
    return out or [para('', size=TABLE_SIZE)]


def weekly_section(xlsx, avail):
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    v = lambda r, c: (str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else '')
    out = [para([(v(1, 1), True)], jc='center', after=60, size=24),
           para(v(2, 1), after=60)]
    r = 3
    while v(r, 2) != 'Даваа':
        a = v(r, 1)
        if a.endswith('СҮД:'):
            out.append(para([(a, True)], keep=True, after=20))
        elif a:
            for ln in a.split('\n'):
                out.append(para(ln, after=20))
        r += 1
    head_row = r
    ncol = 6
    first = 2000
    other = (avail - first) // (ncol - 1)
    widths = [first] + [other] * (ncol - 1)
    tbl = E('tbl')
    tpr = etree.SubElement(tbl, W + 'tblPr')
    tpr.append(E('tblW', w=sum(widths), type='dxa'))
    tpr.append(E('jc', val='center'))
    b = etree.SubElement(tpr, W + 'tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b.append(E(side, val='single', sz=6, space=0, color='000000'))
    tpr.append(E('tblLayout', type='fixed'))
    m = etree.SubElement(tpr, W + 'tblCellMar')
    m.append(E('left', w=70, type='dxa')); m.append(E('right', w=70, type='dxa'))
    grid = etree.SubElement(tbl, W + 'tblGrid')
    for w in widths:
        grid.append(E('gridCol', w=w))

    def tr():
        row = etree.SubElement(tbl, W + 'tr')
        row.append(E('trPr')); row[0].append(E('cantSplit'))   # мөр хуудас дамжиж хуваагдахгүй
        return row

    row = tr()
    for c in range(1, ncol + 1):
        row.append(cell(widths[c - 1], [para([(v(head_row, c), True)], jc='center', after=0, size=TABLE_SIZE)], fill='F2F2F2'))
    merged = {str(x) for x in ws.merged_cells.ranges}
    r = head_row + 1
    while r <= ws.max_row:
        label = v(r, 1)
        pair = f'A{r}:A{r + 1}' in merged
        full = any(x.startswith(f'B{r}:') and x.endswith(f'{r}') and x != f'B{r}:B{r}' for x in merged)
        row = tr()
        row.append(cell(widths[0], [para([(label, True)], after=0, size=TABLE_SIZE)]))
        if full:
            row.append(cell(sum(widths[1:]), bold_title_paras('', v(r, 2)), span=ncol - 1))
        else:
            for c in range(2, ncol + 1):
                row.append(cell(widths[c - 1], bold_title_paras(v(r, c), v(r + 1, c) if pair else '')))
        r += 2 if pair else 1
    out.append(tbl)
    return out


def main(xlsx, daydocx, out):
    z = zipfile.ZipFile(daydocx)
    parts = {n: z.read(n) for n in z.namelist()}
    doc = etree.fromstring(parts['word/document.xml'])
    body = doc.find(W + 'body')
    final = body[-1]
    pw = int(final.find(W + 'pgSz').get(W + 'w'))

    # --- өдрийн хэсэг: фонт, СҮД, хуудаслалт
    list_sud(body)
    kids = list(body)
    for e in kids:
        apply_font(e, True)
    starts = [i for i, e in enumerate(kids) if e.tag == W + 'p' and DAY_START.match(txt(e))]
    # өдөр бүрийн толгой хэсэг: өмнөх хүснэгтийн дараах эхний догол мөрөөс эхэлнэ
    tbls = [i for i, e in enumerate(kids) if e.tag == W + 'tbl']
    for k, s in enumerate(starts):
        b = max([t for t in tbls if t < s], default=-1) + 1
        if k:
            set_flag(kids[b], 'pageBreakBefore')
        for e in kids[b:s + 3]:
            if e.tag == W + 'p':
                set_flag(e, 'keepNext')
    for tbl in body.iter(W + 'tbl'):
        for h in list(tbl.iter(W + 'tblHeader')):
            h.getparent().remove(h)
        for row in tbl.findall(W + 'tr'):
            trpr = row.find(W + 'trPr')
            if trpr is None:
                trpr = E('trPr'); row.insert(0, trpr)
            if trpr.find(W + 'cantSplit') is None:
                trpr.insert(0, E('cantSplit'))
        grid = tbl.find(W + 'tblGrid')
        gw = [int(g.get(W + 'w')) for g in grid]
        avail = pw - LEFT - RIGHT
        if sum(gw) > avail:
            f = (avail - gw[0]) / (sum(gw) - gw[0])
            nw = [gw[0]] + [int(x * f) for x in gw[1:]]
            for g, x in zip(grid, nw):
                g.set(W + 'w', str(x))
            for row in tbl.findall(W + 'tr'):
                col = 0
                for tc in row.findall(W + 'tc'):
                    pr = tc.find(W + 'tcPr'); gs = pr.find(W + 'gridSpan')
                    n = int(gs.get(W + 'val')) if gs is not None else 1
                    tw = pr.find(W + 'tcW')
                    if tw is not None:
                        tw.set(W + 'w', str(sum(nw[col:col + n]))); tw.set(W + 'type', 'dxa')
                    col += n
            tw = tbl.find(f'{W}tblPr/{W}tblW')
            if tw is not None:
                tw.set(W + 'w', str(sum(nw))); tw.set(W + 'type', 'dxa')

    # --- 7 хоногийн хэсэг (хөндлөн) эхэнд
    land = copy.deepcopy(final)
    sz = land.find(W + 'pgSz')
    w, h = int(sz.get(W + 'w')), int(sz.get(W + 'h'))
    sz.set(W + 'w', str(max(w, h))); sz.set(W + 'h', str(min(w, h))); sz.set(W + 'orient', 'landscape')
    weekly = weekly_section(xlsx, max(w, h) - LEFT - RIGHT)
    brk = para('', after=0)
    brk.find(W + 'pPr').append(land)
    for i, e in enumerate(weekly + [brk]):
        body.insert(i, e)

    for sp in body.iter(W + 'sectPr'):
        pm = sp.find(W + 'pgMar')
        for k, val in (('top', TOP), ('left', LEFT), ('bottom', BOTTOM), ('right', RIGHT),
                       ('header', TOP // 2), ('footer', BOTTOM)):
            pm.set(W + k, str(val))

    parts['word/document.xml'] = etree.tostring(doc, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as o:
        for n, d in parts.items():
            o.writestr(n, d)
    print(out)


if __name__ == '__main__':
    main(*sys.argv[1:4])
