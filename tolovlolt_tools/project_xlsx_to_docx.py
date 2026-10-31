"""Төсөлт үйл ажиллагааны Excel төлөвлөлтийг (календарьчилсан, 2 долоо хоног) PROMPT.md-ийн
загвараар долоо хоног бүрт тусдаа .docx болгоно.

Хэрэглээ:
    python3 project_xlsx_to_docx.py ЭХ.xlsx ГАРАХ_ХАВТАС САР ЭХЛЭХ_ДОЛОО_ХОНОГ
    жишээ: python3 project_xlsx_to_docx.py tusul.xlsx out/ 10 3   # 10 сарын 3, 4-р долоо хоног

Бичвэрийг Excel-ээс яг хуулна; зөвхөн хугацааг (сар, долоо хоног) шинэчилнэ.
"""
import os, re, sys
import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT, SIZE = 'Arial', 11
WEEKDAYS = ['Даваа', 'Мягмар', 'Лхагва', 'Пүрэв', 'Баасан']
ROMAN = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V'}
MONTH_ROMAN = {i: r for i, r in enumerate('I II III IV V VI VII VIII IX X XI XII'.split(), 1)}


def ordinal_day(n):
    """1 дэх, 2 дахь, ... (Монгол дагавар)."""
    return f'{n} {"дэх" if n % 10 in (1, 4, 9) else "дахь"}'


def ordinal_month(m):
    return f'{m} {"дүгээр" if m % 10 in (1, 4, 9) or m == 11 else "дугаар"}'


# ---------------------------------------------------------------- Excel унших
def read(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    cell = lambda c, r: (str(ws[f'{c}{r}'].value).strip() if ws[f'{c}{r}'].value is not None else '')
    rows = range(1, ws.max_row + 1)
    find = lambda col, text: next(r for r in rows if cell(col, r).startswith(text))

    d = {'title': cell('A', 1), 'intro': [], 'sud': {}, 'desc': {}, 'cal': [], 'days': [], 'steps': {}}
    r = find('A', 'Төслийн танилцуулга') + 1
    while not cell('A', r).startswith('Төрөл'):
        if cell('A', r):
            d['intro'].append((cell('A', r), cell('B', r)))
        r += 1
    # СҮД: долоо хоног бүрийн Шинэ / Уялдаа / Давталт
    r += 1
    while not cell('A', r).startswith('2 долоо хоногийн') and r < ws.max_row:
        b = cell('B', r)
        if b:
            lines = b.split('\n')
            wk = int(re.search(r'(\d+)-р', lines[0]).group(1))
            for ln in lines[1:]:
                kind, codes = (ln.split(':', 1) if ':' in ln else ('Шинэ', ln))
                d['sud'].setdefault(wk, {}).setdefault(kind.strip(), []).extend(
                    x.strip() for x in codes.split(';') if x.strip())
            for ln in cell('C', r).split('\n'):
                if ':' in ln:
                    code, txt = ln.split(':', 1)
                    d['desc'][code.strip()] = ln.strip()
        r += 1
    # хураангуй календарь
    r = find('A', 'Долоо хоног') + 1
    while cell('A', r) and 'долоо хоног' in cell('A', r):
        d['cal'].append([cell(c, r) for c in 'ABCDEFGH'])
        r += 1
    d['cal_head'] = [cell(c, find('A', 'Долоо хоног')) for c in 'ABCDEFGH']
    # өдөр бүрийн дэлгэрэнгүй
    r = next(x for x in rows if cell('B', x).startswith('Үйл ажиллагааны нэр'))
    d['day_head'] = [cell(c, r) for c in 'ABCDEFGH']
    r += 1
    while cell('A', r) and 'Өдөр' in cell('A', r):
        d['days'].append([cell(c, r) for c in 'ABCDEFGH'])
        r += 1
    # ерөнхий төлөвлөлт (алхам)
    r = find('A', 'Ерөнхий төлөвлөлт') + 1
    d['gen_head'] = [cell('A', r), cell('B', r)]
    step = None
    for r in range(r + 1, ws.max_row + 1):
        if cell('A', r).startswith('Алхам'):
            step = cell('A', r)
            d['gen_cols'] = [cell('F', r), cell('G', r), cell('H', r)]
        elif cell('A', r):
            n = int(re.match(r'\d+', cell('H', r)).group())
            d['steps'][n] = dict(step=step, codes=cell('A', r), content=cell('B', r),
                                 name=cell('F', r), crit=cell('G', r), when=cell('H', r))
    return d


# ---------------------------------------------------------------- Word туслах
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def style_run(run, bold=False, size=SIZE):
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)


def para(container, parts, align=None, after=2, keep=False, page_break=False, size=SIZE):
    """parts: [(бичвэр, тод эсэх), ...] эсвэл нэг мөр бичвэр."""
    p = container.add_paragraph()
    if isinstance(parts, str):
        parts = [(parts, False)]
    for k, (t, b) in enumerate(parts):
        lines = t.split('\n')
        for j, ln in enumerate(lines):
            r = p.add_run(ln)
            style_run(r, b, size)
            if j < len(lines) - 1:
                r.add_break()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(after)
    pf.line_spacing = 1.08
    if align:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if page_break:
        pf.page_break_before = True
    return p


def labeled(container, line, **kw):
    """«Шошго: бичвэр» → шошгыг тод."""
    m = re.match(r'^([^:]{1,60}):\s*(.*)$', line, re.S)
    if m:
        return para(container, [(m.group(1) + ': ', True), (m.group(2), False)], **kw)
    return para(container, line, **kw)


def sud_block(container, title, codes, desc, keep=False):
    para(container, [(title, True)], keep=True)
    for c in codes:
        para(container, desc.get(c, c), keep=keep)


def table(doc, head, rows, widths, merged_rows=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate([head] + rows):
        tr = t.rows[i]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))          # мөр хуудас дамжиж хуваагдахгүй
        if i in merged_rows:
            c = t.rows[i].cells[0].merge(t.rows[i].cells[-1])
            c.text = ''
            p = c.paragraphs[0]
            style_run(p.add_run(row[0]), True)
            set_cell_shading(c, 'F2F2F2')
            continue
        for j, val in enumerate(row):
            c = t.rows[i].cells[j]
            c.width = widths[j]
            p = c.paragraphs[0]
            for k, ln in enumerate(val.split('\n')):
                if k:
                    p = c.add_paragraph()
                m = re.match(r'^([^:]{1,40}):\s*(.*)$', ln) if i else None
                if m and m.group(2):
                    style_run(p.add_run(m.group(1) + ': '), True)
                    style_run(p.add_run(m.group(2)))
                else:
                    style_run(p.add_run(ln), i == 0)
                p.paragraph_format.space_after = Pt(0)
            if i == 0:
                set_cell_shading(c, 'F2F2F2')
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j, w in enumerate(widths):
        t.columns[j].width = w
    return t


def margins(sec, landscape):
    sec.page_width, sec.page_height = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.top_margin, sec.left_margin = Cm(1.5), Cm(1.5)
    sec.bottom_margin, sec.right_margin = Cm(0.5), Cm(0.5)
    sec.header_distance, sec.footer_distance = Cm(0.75), Cm(0.5)


# ---------------------------------------------------------------- баримт үүсгэх
def build(d, wk_idx, month, week_no, outdir, file_no):
    days = d['days'][wk_idx * 5:(wk_idx + 1) * 5]
    src_wk = wk_idx + 1                                   # Excel дээрх 1-р, 2-р долoo хоног
    sud = d['sud'][src_wk]
    week_label = f'{month} сарын {week_no}-р долоо хоног'
    topic = next(v for k, v in d['intro'] if k == 'Төслийн нэр')
    group = next(v for k, v in d['intro'] if k == 'Бүлэг')
    group_gen = group.split()[0].upper() + ' БҮЛГИЙН'

    doc = Document()
    st = doc.styles['Normal']
    st.font.name, st.font.size = FONT, Pt(SIZE)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    zoom = doc.settings.element.find(qn('w:zoom'))
    if zoom is not None and zoom.get(qn('w:percent')) is None:
        zoom.set(qn('w:percent'), '100')

    # 1) Төслийн танилцуулга (босоо)
    margins(doc.sections[0], False)
    para(doc, [(d['title'], True)], align=WD_ALIGN_PARAGRAPH.CENTER, after=8, size=13)
    para(doc, [('Төслийн танилцуулга', True)], after=4, size=12)
    first, last = week_no, week_no + len(d['days']) // 5 - 1
    for k, v in d['intro']:
        if k == 'Хугацаа':
            v = re.sub(r'^\d+ сарын \d+[–-]\d+-р', f'{month} сарын {first}–{last}-р', v)
        labeled(doc, f'{k}: {v}', after=4)

    # 2) 7 хоногийн төлөвлөлт (хөндлөн)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    margins(sec, True)
    para(doc, [('БАТЛАВ. АРГА ЗҮЙЧ: .................................', True)], align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, [(f'{group_gen} ТӨСӨЛТ ҮЙЛ АЖИЛЛАГААНЫ 7 ХОНОГИЙН ТӨЛӨВЛӨЛТ', True)],
         align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, [('Сэдэв: ', True), (f'{topic}. {week_label}.', False)])
    para(doc, [('Суралцахуйн үр дүн:', True)], keep=True)
    for kind in ('Шинэ', 'Уялдаа', 'Давталт'):
        if sud.get(kind):
            sud_block(doc, f'{kind} СҮД:', sud[kind], d['desc'])
    para(doc, '', after=0)
    cal = [r[:] for r in d['cal'] if r[0].startswith(f'{src_wk}-р')]
    for r in cal:
        r[0] = week_label
    W = Cm(27.7)
    table(doc, d['cal_head'], cal, [Cm(2.6)] + [Cm(3.6)] * 5 + [Cm(3.6), Cm(3.5)])
    para(doc, '', after=0)
    rows, merged = [], []
    step = None
    for i, _ in enumerate(days):
        n = wk_idx * 5 + i + 1
        g = d['steps'][n]
        if g['step'] != step:
            step = g['step']
            merged.append(len(rows) + 1)
            rows.append([step, '', '', '', ''])
        rows.append([g['codes'], g['content'], g['name'], g['crit'], g['when']])
    table(doc, [d['gen_head'][0], d['gen_head'][1]] + d['gen_cols'], rows,
          [Cm(3.2), Cm(9.0), Cm(4.2), Cm(8.8), Cm(2.5)], merged_rows=merged)
    para(doc, 'Төлөвлөгөө боловсруулсан багш: .................................', after=0).paragraph_format.space_before = Pt(4)

    # 3) Өдөр бүрийн төлөвлөлт (босоо, өдөр бүр шинэ хуудаснаас)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    margins(sec, False)
    para(doc, [(f'{ordinal_month(month)} сарын {ROMAN[week_no]} долоо хоног', True)],
         align=WD_ALIGN_PARAGRAPH.CENTER, after=6, keep=True, size=12)
    all_codes = {c: k for k in ('Шинэ', 'Уялдаа', 'Давталт') for c in sud.get(k, [])}
    for i, row in enumerate(days):
        n = wk_idx * 5 + i + 1
        g = d['steps'][n]
        a_lines, b_lines = row[0].split('\n'), row[1].split('\n')
        name, codes = b_lines[0], [c.strip() for c in b_lines[1].split(';')]
        question = '\n'.join(b_lines[2:])
        para(doc, [(f'{name.upper()} ({ordinal_day(n)} өдөр - {WEEKDAYS[i]} гараг)', True)],
             after=4, keep=True, page_break=i > 0)
        para(doc, [('Сэдэв: ', True), (f'{topic}. Дэд сэдэв: {name}', False)], keep=True)
        para(doc, [('Төслийн үе шат: ', True), (g['step'], False)], keep=True)
        for kind in ('Шинэ', 'Уялдаа', 'Давталт'):
            mine = [c for c in codes if all_codes.get(c) == kind]
            if mine:
                sud_block(doc, f'{kind} СҮД:', mine, d['desc'], keep=True)
        labeled(doc, question, after=4)
        labeled(doc, f'{d["gen_head"][1]}: {g["content"]}', after=4)
        for col in (2, 3, 4, 5):
            for ln in row[col].split('\n'):
                labeled(doc, ln, after=4)
        labeled(doc, f'{d["day_head"][6]}: {row[6]}', after=4)
        labeled(doc, f'{d["gen_cols"][1]}: {g["crit"]}', after=4)
        labeled(doc, f'{d["day_head"][7]}: {row[7]}', after=4)

    fn = os.path.join(outdir, f'{file_no}. {topic.upper()} ({group.upper()}) - '
                              f'{MONTH_ROMAN[month]} сарын {ROMAN[week_no]} долоо хоног.docx')
    doc.save(fn)
    print(fn)


if __name__ == '__main__':
    src, outdir, month, week = sys.argv[1], sys.argv[2], int(sys.argv[3]), int(sys.argv[4])
    data = read(src)
    for w in range(len(data['days']) // 5):
        build(data, w, month, week + w, outdir, w + 1)
